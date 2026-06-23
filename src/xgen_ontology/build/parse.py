"""Light document parsing — extract text from common file formats.

Built-in (zero deps): ``.txt`` ``.md`` ``.rst`` ``.json`` ``.log`` (raw),
``.html``/``.htm`` (stdlib tag strip), and ``.csv``/``.tsv`` (kept as raw text so
the tabular builder can parse the table). Optional via the ``[files]`` extra:
``.pdf`` (pypdf), ``.docx`` (python-docx), ``.xlsx`` (openpyxl -> CSV text).
"""
from __future__ import annotations

import os
from html.parser import HTMLParser

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".json", ".log", ".text", ""}
TABLE_EXT = {".csv", ".tsv"}
HTML_EXT = {".html", ".htm", ".xhtml"}


def extract_text(path: str) -> str:
    """Extract plain text (or raw table text) from a file by extension."""
    ext = os.path.splitext(path)[1].lower()
    if ext in HTML_EXT:
        return html_to_text(_read_bytes(path).decode("utf-8", "ignore"))
    if ext == ".pdf":
        return _pdf_to_text(path)
    if ext == ".docx":
        return _docx_to_text(path)
    if ext == ".xlsx":
        return _xlsx_to_csv(path)
    # text + csv/tsv: read as text (csv/tsv kept raw for the tabular builder)
    return _read_bytes(path).decode("utf-8-sig", "ignore")


def load_documents(paths: list[str]) -> dict[str, str]:
    """``{basename: text}`` for a list of files — feed straight to build_from_documents."""
    return {os.path.basename(p): extract_text(p) for p in paths}


def _read_bytes(path: str) -> bytes:
    with open(path, "rb") as f:
        return f.read()


class _HTMLText(HTMLParser):
    _BLOCK = {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip and data.strip():
            self.parts.append(data)


def html_to_text(html: str) -> str:
    p = _HTMLText()
    try:
        p.feed(html)
    except Exception:
        return html
    lines = [ln.strip() for ln in "".join(p.parts).splitlines()]
    return "\n".join(ln for ln in lines if ln)


def _pdf_to_text(path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("PDF parsing needs pypdf: pip install 'xgen-ontology[files]'") from e
    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx_to_text(path: str) -> str:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("DOCX parsing needs python-docx: pip install 'xgen-ontology[files]'") from e
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs if p.text.strip())


def _xlsx_to_csv(path: str) -> str:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("XLSX parsing needs openpyxl: pip install 'xgen-ontology[files]'") from e
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(",".join("" if c is None else str(c) for c in row))
    return "\n".join(rows)
